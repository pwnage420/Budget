"""Tests for the Excel + PDF + CSV exporters.

Confirms structure (sheet names, page count, columns) so CI catches regressions
without running the full demo.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest

from app.core.sample_data import generate_price_series
from app.core.data_loader import load_series_from_dataframe
from app.export.csv_writer import write_budget_wide, write_history_wide
from app.export.excel_writer import ExportPayload, _sanitize, export_workbook
from app.export.pdf_writer import PdfPayload, export_pdf

EXPECTED_SHEETS = [
    "README", "Glossary", "Dashboard", "Budget 2026", "FY Tracking",
    "Source Data", "Model Comparison", "Methodology", "Backtest",
]


@pytest.fixture
def synthetic_payload():
    history = load_series_from_dataframe(generate_price_series(seed=42)).df
    forecast = pd.DataFrame({
        "date": pd.date_range("2026-01-05", periods=52, freq="W-MON"),
        "mean": [3.0 + 0.1 * (i % 5) for i in range(52)],
        "p20": [2.7 + 0.1 * (i % 5) for i in range(52)],
        "p30": [2.8 + 0.1 * (i % 5) for i in range(52)],
        "p70": [3.2 + 0.1 * (i % 5) for i in range(52)],
        "p80": [3.3 + 0.1 * (i % 5) for i in range(52)],
    })
    metrics = pd.DataFrame([
        {"model": "Naive", "rmse": 0.3, "mae": 0.2, "mape": 0.1, "smape": 0.1,
         "mase": 1.0, "bias": 0.01, "coverage_p20_p80": 0.6, "fit_seconds": 0.1,
         "n_folds": 1, "failed": False},
    ])
    return history, forecast, metrics


def test_excel_all_sheets_present(tmp_path, synthetic_payload):
    history, forecast, metrics = synthetic_payload
    payload = ExportPayload(
        history=history, forecast=forecast, metrics=metrics,
        blend_weights={"Naive": 1.0}, blend_strategy="best_single",
        blend_backtest_rmse=0.3,
    )
    path = export_workbook(payload, tmp_path / "test.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    for sheet in EXPECTED_SHEETS:
        assert sheet in wb.sheetnames, f"Missing sheet: {sheet}"
    assert wb.sheetnames == EXPECTED_SHEETS, "Sheet order must match spec"


def test_excel_injection_guard():
    """Strings starting with =, +, -, @ must be prefixed with a single quote."""
    assert _sanitize("=CMD|'/C calc'!A1") == "'=CMD|'/C calc'!A1"
    assert _sanitize("+HYPERLINK(...)") == "'+HYPERLINK(...)"
    assert _sanitize("-1") == "'-1"
    assert _sanitize("@SUM(A:A)") == "'@SUM(A:A)"
    assert _sanitize("benign") == "benign"
    assert _sanitize(3.14) == 3.14


def test_excel_handles_pandas_na(tmp_path, synthetic_payload):
    """pd.NA / pd.NaT / NaN must coerce to None for openpyxl."""
    history, forecast, metrics = synthetic_payload
    actuals = pd.DataFrame({
        "date": [pd.Timestamp("2026-01-05"), pd.Timestamp("2026-01-12")],
        "actual_price": [3.10, 3.20],
    })
    payload = ExportPayload(
        history=history, forecast=forecast, metrics=metrics,
        blend_weights={"Naive": 1.0}, blend_strategy="best_single",
        blend_backtest_rmse=0.3, actuals=actuals,
    )
    path = export_workbook(payload, tmp_path / "test.xlsx")
    assert path.exists()


def test_pdf_has_four_pages(tmp_path, synthetic_payload):
    history, forecast, metrics = synthetic_payload
    payload = PdfPayload(
        history=history, forecast=forecast, actuals=None, metrics=metrics,
        blend_weights={"Naive": 1.0}, blend_strategy="best_single",
        blend_backtest_rmse=0.3,
    )
    path = export_pdf(payload, tmp_path / "test.pdf")
    raw = path.read_bytes()
    # Count /Type /Page occurrences (not /Type /Pages)
    n_page = (raw.count(b"/Type /Page\n") + raw.count(b"/Type /Page ")
               + raw.count(b"/Type/Page\n") + raw.count(b"/Type/Page "))
    assert n_page == 4, f"Expected 4 PDF pages, got {n_page}"


def test_budget_wide_csv(tmp_path, synthetic_payload):
    _, forecast, _ = synthetic_payload
    path = write_budget_wide(forecast, tmp_path / "budget.csv")
    df = pd.read_csv(path)
    assert list(df.columns) == ["Year", "Week", "P20", "P50", "P80"]
    assert len(df) == 52
    # P20 <= P50 <= P80 row-wise
    assert (df["P20"] <= df["P50"]).all()
    assert (df["P50"] <= df["P80"]).all()


def test_history_wide_csv_roundtrips(tmp_path, synthetic_payload):
    history, _, _ = synthetic_payload
    path = write_history_wide(history, tmp_path / "history.csv")
    df = pd.read_csv(path)
    assert df.columns[0] == "Week"
    assert "2021" in df.columns
    assert "2025" in df.columns
    assert len(df) == 52


def test_budget_wide_columns_configurable(tmp_path, synthetic_payload):
    _, forecast, _ = synthetic_payload
    path = write_budget_wide(forecast, tmp_path / "budget_p50only.csv",
                              columns=("P50",))
    df = pd.read_csv(path)
    assert list(df.columns) == ["Year", "Week", "P50"]
